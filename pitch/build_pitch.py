"""Genera el pitch de Pixel Pattern en PowerPoint.

Uso (desde la raíz del repo, en Windows):

    backend\\.venv\\Scripts\\python.exe pitch\\build_pitch.py

Lee las capturas de ``pitch/img/``. Si falta alguna, dibuja un recuadro gris
con el texto "Captura: <nombre>" para sustituirla más tarde. Al volver a
ejecutar el script con las capturas ya en su sitio, aparecen solas.

Requisito: python-pptx instalado en el venv (no está en requirements.txt).
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Sistema de diseño
# ---------------------------------------------------------------------------

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.85
CONTENT_W = SLIDE_W - 2 * MARGIN

BG = RGBColor.from_string("FBF8F4")
INK = RGBColor.from_string("1F1D2B")
ACCENT = RGBColor.from_string("E8456B")
ACCENT_DARK = RGBColor.from_string("CB3155")
ACCENT_LIGHT = RGBColor.from_string("F28CA4")
ACCENT_SOFT = RGBColor.from_string("FBE7ED")
WHITE = RGBColor.from_string("FFFFFF")
MUTED = RGBColor.from_string("6E6A7C")
BORDER = RGBColor.from_string("E7DED4")
PLACEHOLDER = RGBColor.from_string("ECE5DC")
PLACEHOLDER_LINE = RGBColor.from_string("C9BFB4")

FONT = "Segoe UI"
TOTAL_SLIDES = 11

IMG_DIR = Path(__file__).resolve().parent / "img"
OUT_PATH = Path(__file__).resolve().parent / "pixel-pattern-pitch.pptx"


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------


def set_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _style_run(run, size: float, color: RGBColor, bold: bool, italic: bool,
               tracking: float | None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if tracking:
        rpr = run._r.get_or_add_rPr()
        rpr.set("spc", str(int(tracking * 100)))


def add_text(slide, text: str, x: float, y: float, w: float, h: float,
             size: float = 18, color: RGBColor = INK, bold: bool = False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing: float = 1.0, space_after: float = 0,
             italic: bool = False, tracking: float | None = None,
             margins: tuple[float, float, float, float] = (0, 0, 0, 0)):
    """Añade una caja de texto. Respeta los saltos de línea como párrafos."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (
        Inches(margins[0]), Inches(margins[1]), Inches(margins[2]), Inches(margins[3])
    )

    paragraphs = text.split("\n")
    for i, line in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(space_after if i < len(paragraphs) - 1 else 0)
        run = p.add_run()
        run.text = line
        _style_run(run, size, color, bold, italic, tracking)
    return box


def add_rich_text(slide, runs, x: float, y: float, w: float, h: float,
                  size: float = 18, align=PP_ALIGN.LEFT,
                  anchor=MSO_ANCHOR.MIDDLE, line_spacing: float = 1.0):
    """Caja de texto con varios tramos con estilo propio (lista de dicts)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for spec in runs:
        run = p.add_run()
        run.text = spec["t"]
        _style_run(
            run,
            spec.get("size", size),
            spec.get("color", INK),
            spec.get("bold", False),
            spec.get("italic", False),
            spec.get("tracking"),
        )
    return box


def soft_shadow(shape, blur: int = 152400, dist: int = 30000,
                alpha: int = 9000, direction: int = 5400000) -> None:
    """Sombra exterior suave vía XML (python-pptx no la expone)."""
    effect = (
        f'<a:effectLst {nsdecls("a")}>'
        f'<a:outerShdw blurRad="{blur}" dist="{dist}" dir="{direction}" '
        'algn="ctr" rotWithShape="0">'
        f'<a:srgbClr val="1F1D2B"><a:alpha val="{alpha}"/></a:srgbClr>'
        "</a:outerShdw></a:effectLst>"
    )
    shape._element.spPr.append(parse_xml(effect))


def _strip_theme_style(shape) -> None:
    """Quita el estilo del tema (incluida su sombra) para control total."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def add_rect(slide, x: float, y: float, w: float, h: float,
             fill: RGBColor | None = None, line: RGBColor | None = None,
             line_w: float = 1.0, radius: float | None = None,
             shadow: bool = False, dash: MSO_LINE_DASH_STYLE | None = None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    _strip_theme_style(shape)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
        if dash is not None:
            shape.line.dash_style = dash
    if radius is not None:
        shape.adjustments[0] = radius
    if shadow:
        soft_shadow(shape)
    return shape


def add_arrow(slide, x: float, y: float, w: float, h: float,
              fill: RGBColor = RGBColor.from_string("C9C0D2")) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    _strip_theme_style(shape)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_eyebrow(slide, text: str, color: RGBColor = ACCENT) -> None:
    add_text(slide, text.upper(), MARGIN, 0.5, 8.0, 0.3, size=10.5,
             color=color, bold=True, tracking=1.9)


def add_title(slide, text: str, size: float = 40) -> None:
    add_text(slide, text, MARGIN, 0.84, CONTENT_W, 1.15, size=size,
             color=INK, bold=True, line_spacing=0.95)


def add_rule(slide) -> None:
    add_rect(slide, MARGIN, 2.03, 0.62, 0.055, fill=ACCENT)


def add_footer(slide, number: int) -> None:
    add_rect(slide, MARGIN, 6.92, CONTENT_W, 0.009, fill=BORDER)
    add_text(slide, "Pixel Pattern", MARGIN, 7.02, 5.0, 0.25, size=9,
             color=MUTED)
    add_text(slide, f"{number:02d} / {TOTAL_SLIDES}", 10.5, 7.02, 1.98, 0.25,
             size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def image_size(path: Path) -> tuple[int, int] | None:
    try:
        from PIL import Image

        with Image.open(path) as im:
            return im.width, im.height
    except Exception:
        return None


def add_screenshot(slide, filename: str, x: float, y: float,
                   w: float, h: float) -> None:
    """Tarjeta blanca con la captura dentro o un placeholder gris si falta."""
    card = add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER,
                    line_w=1.0, radius=0.035, shadow=True)
    pad = 0.13
    path = IMG_DIR / filename

    if path.exists():
        dims = image_size(path)
        avail_w, avail_h = w - 2 * pad, h - 2 * pad
        if dims:
            iw, ih = dims
            scale = min(avail_w / iw, avail_h / ih)
            pw, ph = iw * scale, ih * scale
        else:
            pw, ph = avail_w, avail_h
        px = x + (w - pw) / 2
        py = y + (h - ph) / 2
        pic = slide.shapes.add_picture(
            str(path), Inches(px), Inches(py), Inches(pw), Inches(ph)
        )
        pic.line.color.rgb = BORDER
        pic.line.width = Pt(0.75)
    else:
        add_rect(slide, x + pad, y + pad, w - 2 * pad, h - 2 * pad,
                 fill=PLACEHOLDER, line=PLACEHOLDER_LINE, line_w=1.25,
                 radius=0.02, dash=MSO_LINE_DASH_STYLE.DASH)
        add_text(
            slide,
            f"Captura: {filename}",
            x + pad,
            y + h / 2 - 0.2,
            w - 2 * pad,
            0.4,
            size=14,
            color=MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    return card


def add_pixel_mark(slide, x: float, y: float, cell: float = 0.12,
                   gap: float = 0.035, colors: tuple[RGBColor, ...] = (ACCENT, ACCENT_LIGHT, INK, ACCENT_DARK)) -> None:
    """Pequeño motivo de 2x2 píxeles, marca de la casa."""
    coords = [(0, 0), (1, 0), (0, 1), (1, 1)]
    for (cx, cy), color in zip(coords, colors):
        add_rect(slide, x + cx * (cell + gap), y + cy * (cell + gap),
                 cell, cell, fill=color)


def add_heart(slide, x: float, y: float, cell: float = 0.36,
              gap: float = 0.04) -> None:
    """Corazón pixelado, el motivo visual de la portada."""
    heart = [
        ".XX...XX.",
        "XXXX.XXXX",
        "XXXXXXXXX",
        "XXXXXXXXX",
        ".XXXXXXX.",
        "..XXXXX..",
        "...XXX...",
        "....X....",
    ]
    for row, line in enumerate(heart):
        for col, char in enumerate(line):
            if char != "X":
                continue
            add_rect(
                slide,
                x + col * (cell + gap),
                y + row * (cell + gap),
                cell,
                cell,
                fill=ACCENT,
            )


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------
# Diapositivas
# ---------------------------------------------------------------------------


def slide_cover(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)

    add_eyebrow(s, "Hackathon · [Nombre del hackathon]")
    add_text(s, "Pixel\nPattern", MARGIN, 1.28, 7.2, 2.45, size=68,
             color=INK, bold=True, line_spacing=0.88)
    add_rect(s, MARGIN + 0.02, 3.93, 1.15, 0.07, fill=ACCENT)
    add_text(s, "Crea, comparte y haz patrones\nde HAMA y punto de cruz",
             MARGIN + 0.02, 4.3, 6.6, 1.1, size=22, color=MUTED,
             line_spacing=1.18)

    add_heart(s, 8.92, 1.58, cell=0.365, gap=0.04)
    add_text(s, "1 píxel = 1 cuenta", 8.92, 5.18, 3.65, 0.3, size=11.5,
             color=MUTED, align=PP_ALIGN.CENTER, tracking=1.2)

    add_text(s, "[Nombre del hackathon]", MARGIN, 6.08, 6.5, 0.3,
             size=13, color=INK, bold=True)
    add_text(s, "[Nombre del equipo] · 2026", MARGIN, 6.4, 6.5, 0.3,
             size=13, color=MUTED)

    add_notes(
        s,
        "≈ 15 s — Buenas, somos [Nombre del equipo] y esto es Pixel Pattern. "
        "En cuatro minutos os enseñamos cómo convertir una foto en un patrón de "
        "HAMA o punto de cruz y cómo compartirlo con una comunidad que los hace. "
        "Vamos al lío.",
    )


def slide_problem(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "El problema")
    add_title(s, "Crear un patrón desde una foto es tedioso")
    add_rule(s)

    rows = [
        "Contar cuentas a mano",
        "Colores que no existen en tu marca",
        "Patrones dispersos, sin progreso",
    ]
    for i, text in enumerate(rows):
        y = 2.55 + i * 0.92
        add_text(s, f"0{i + 1}", MARGIN, y + 0.02, 0.55, 0.35, size=11.5,
                 color=ACCENT, bold=True, tracking=1.2)
        add_text(s, text, MARGIN + 0.72, y - 0.02, 5.9, 0.55, size=19,
                 color=INK)
        if i < 2:
            add_rect(s, MARGIN + 0.72, y + 0.62, 5.15, 0.009, fill=BORDER)

    card_x, card_y, card_w, card_h = 8.18, 2.45, 4.3, 3.6
    add_rect(s, card_x, card_y, card_w, card_h, fill=WHITE, line=BORDER,
             radius=0.04, shadow=True)
    add_text(s, "4.800", card_x + 0.36, card_y + 0.3, card_w - 0.7, 1.05,
             size=62, color=ACCENT, bold=True, line_spacing=0.9)
    add_text(s, "cuentas en 80 × 60", card_x + 0.36, card_y + 1.35,
             card_w - 0.7, 0.3, size=14, color=MUTED)

    # Rejilla decorativa: un patrón a medio hacer.
    gx, gy, cell, gap = card_x + 0.38, card_y + 2.02, 0.155, 0.045
    accent_cells = {(0, 1), (0, 4), (1, 2), (2, 0), (2, 5), (3, 3), (4, 1), (4, 4)}
    for r in range(5):
        for c in range(8):
            color = ACCENT if (r, c) in accent_cells else BORDER
            add_rect(s, gx + c * (cell + gap), gy + r * (cell + gap),
                     cell, cell, fill=color)

    add_footer(s, 2)
    add_notes(
        s,
        "≈ 30 s — Si alguna vez habéis hecho un patrón desde una foto, conocéis "
        "el dolor: cuentas cuadraditos a mano y buscas un color que se parezca... "
        "y resulta que no existe en Hama ni en DMC. Y lo que hay en Pinterest o "
        "en blogs está en una imagen, sin lista de materiales ni progreso. El "
        "patrón es bonito; hacerlo es otra historia.",
    )


def slide_solution(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "La solución")
    add_title(s, "Una app, todo el ciclo")
    add_rule(s)

    cards = [
        ("CREA", "Foto o dibujo → patrón con colores Hama y DMC."),
        ("HAZ", "Edita, exporta e imprime."),
        ("COMPARTE", "Feed, perfiles y remix."),
    ]
    for i, (label, body) in enumerate(cards):
        y = 2.5 + i * 1.22
        add_pixel_mark(s, MARGIN + 0.02, y + 0.02, cell=0.13, gap=0.035)
        add_text(s, label, MARGIN + 0.5, y, 2.2, 0.3, size=10.5,
                 color=ACCENT, bold=True, tracking=1.9)
        add_text(s, body, MARGIN + 0.5, y + 0.32, 4.6, 0.78, size=18,
                 color=INK, line_spacing=1.18)

    add_screenshot(s, "crear.png", 6.95, 2.2, 5.5, 3.95)

    add_footer(s, 3)
    add_notes(
        s,
        "≈ 25 s — Pixel Pattern junta las tres piezas que faltaban. Conviertes "
        "una foto o un dibujo en un patrón con colores reales de fabricante; lo "
        "editas cuenta a cuenta y te llevas un PDF con la lista de materiales; y "
        "lo publicas en una comunidad con feed, perfiles y remix. Crear, hacer y "
        "compartir, por fin juntos.",
    )


def slide_demo_convert(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "Demo · 01")
    add_title(s, "De foto a patrón")
    add_rule(s)

    steps = [
        "Sube tu foto",
        "Simplifica a dibujo",
        "Colores reales Hama / DMC",
        "Retoca cuenta a cuenta",
    ]
    rail_x = 1.16
    rail_top = 2.76
    rail_bottom = 5.62
    add_rect(s, rail_x - 0.006, rail_top, 0.012, rail_bottom - rail_top,
             fill=BORDER)
    for i, text in enumerate(steps):
        cy = 2.55 + i * 0.95
        add_rect(s, rail_x - 0.21, cy, 0.42, 0.42, fill=ACCENT_SOFT,
                 line=None, radius=0.5)
        add_text(s, f"{i + 1:02d}", rail_x - 0.21, cy, 0.42, 0.42, size=10.5,
                 color=ACCENT, bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, text, 1.78, cy - 0.05, 4.25, 0.52, size=18, color=INK,
                 anchor=MSO_ANCHOR.MIDDLE)

    add_screenshot(s, "foto-a-patron.png", 6.33, 1.52, 6.17, 4.98)
    add_footer(s, 4)
    add_notes(
        s,
        "≈ 30 s — Aquí está el corazón del producto. Subes una foto y, con un "
        "par de ajustes, la simplificamos a un dibujo: menos colores y formas más "
        "limpias. Después la pasamos a la paleta real de Hama o DMC, así que cada "
        "color se puede comprar. Y si algo no te convence, cambias las cuentas "
        "una a una en el editor.",
    )


def slide_engine(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "Cómo funciona")
    add_title(s, "Del píxel al patrón")
    add_rule(s)

    steps = [
        ("Foto", ""),
        ("Suavizado", "OpenCV"),
        ("Rejilla", ""),
        ("k-means", "CIELAB"),
        ("Paleta real", "Hama · DMC"),
        ("Limpieza", "píxeles sueltos"),
        ("Patrón", ""),
    ]
    cw, ch = 1.4, 1.8
    gap = (CONTENT_W - 7 * cw) / 6
    for i, (label, sub) in enumerate(steps):
        x = MARGIN + i * (cw + gap)
        filled = i in (0, 3, 6)
        if i == 3:
            fill, line, fg, sub_fg = ACCENT, None, WHITE, WHITE
        elif filled:
            fill, line, fg, sub_fg = INK, None, WHITE, BORDER
        else:
            fill, line, fg, sub_fg = WHITE, BORDER, INK, MUTED

        add_rect(s, x, 2.55, cw, ch, fill=fill, line=line, radius=0.05,
                 shadow=not filled)
        add_text(s, f"{i + 1:02d}", x + 0.13, 2.68, 0.5, 0.22, size=8.5,
                 color=sub_fg, bold=True, tracking=0.8)
        add_text(s, label, x + 0.1, 3.0, cw - 0.2, 0.62, size=12.5,
                 color=fg, bold=True, align=PP_ALIGN.CENTER,
                 line_spacing=0.95)
        add_text(s, sub, x + 0.08, 3.76, cw - 0.16, 0.38, size=8.5,
                 color=sub_fg, align=PP_ALIGN.CENTER, line_spacing=1.0)

        if i < len(steps) - 1:
            add_arrow(s, x + cw + 0.064, 3.36, 0.18, 0.18)

    add_rect(s, MARGIN, 4.85, CONTENT_W, 1.0, fill=ACCENT_SOFT, radius=0.05)
    add_rich_text(
        s,
        [
            {"t": "CIELAB", "bold": True, "color": ACCENT},
            {"t": " — distancia de color como la percibe el ojo.",
             "color": INK},
        ],
        MARGIN + 0.36, 4.85, CONTENT_W - 0.72, 1.0, size=16.5,
    )

    add_footer(s, 5)
    add_notes(
        s,
        "≈ 40 s — Por dentro es una cadena de pasos. Suavizamos con OpenCV para "
        "quitar ruido; reducimos a la rejilla; agrupamos los colores con k-means "
        "en espacio CIELAB, no en RGB. ¿Por qué? Porque CIELAB mide la distancia "
        "de color como la percibe el ojo: dos rojos casi iguales quedan cerca, y "
        "no nos llevamos un azul raro solo porque lo digan los números. Después "
        "emparejamos con la paleta real del fabricante, limpiamos píxeles sueltos "
        "y listo.",
    )


def slide_demo_make(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "Demo · 02")
    add_title(s, "Hazlo tú, cuenta a cuenta")
    add_rule(s)

    add_screenshot(s, "hacer.png", MARGIN, 1.78, 6.17, 4.68)

    rows = [
        "Modo por color",
        "Progreso cuenta a cuenta",
        "Lista de materiales y placas",
        "PDF imprimible",
    ]
    for i, text in enumerate(rows):
        y = 2.55 + i * 0.86
        add_rect(s, 7.38, y + 0.16, 0.13, 0.13, fill=ACCENT)
        add_text(s, text, 7.72, y, 4.6, 0.52, size=19, color=INK,
                 anchor=MSO_ANCHOR.MIDDLE)

    add_footer(s, 6)
    add_notes(
        s,
        "≈ 30 s — Una vez publicado, cualquiera puede abrirlo en modo hacer. Se "
        "sigue color por color, que es como se trabaja de verdad: pones todas las "
        "cuentas de un color y no vas cambiando de bote. Marcas las que llevas, ves "
        "el progreso y tienes la lista de materiales con las cuentas y las placas. "
        "Al terminar, exportas el PDF y lo imprimes.",
    )


def slide_community(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "Comunidad")
    add_title(s, "Comparte. Descubre. Remezcla.")
    add_rule(s)

    add_screenshot(s, "feed.png", MARGIN, 2.3, 5.68, 3.44)
    add_screenshot(s, "buscador.png", 6.81, 2.3, 5.68, 3.44)

    pills = [
        ("Feed tipo Instagram", 2.05),
        ("Likes, guardados, comentarios", 2.98),
        ("Seguir y remix", 1.86),
        ("Buscador: perfiles, patrones y #tags", 4.15),
    ]
    x = MARGIN
    for text, width in pills:
        add_rect(s, x, 6.02, width, 0.44, fill=WHITE, line=BORDER,
                 radius=0.5)
        add_text(s, text, x, 6.02, width, 0.44, size=10.5, color=MUTED,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += width + 0.18

    add_footer(s, 7)
    add_notes(
        s,
        "≈ 30 s — La parte social es intencionadamente familiar: feed tipo "
        "Instagram, con likes, guardados y comentarios; perfiles a los que seguir; "
        "y un botón de remix para adaptar el patrón de otra persona. El buscador "
        "encuentra perfiles, patrones y hashtags: buscas gato o pixel art y "
        "encuentras ideas y gente a la que seguir. La comunidad es lo que hace "
        "crecer el catálogo.",
    )


def slide_tech(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "Tecnología")
    add_title(s, "Simple por fuera, sólido por dentro")
    add_rule(s)

    boxes = [
        {
            "label": "FRONTEND",
            "tech": "React · TypeScript · Tailwind",
            "fill": WHITE, "line": BORDER, "fg": INK, "label_color": ACCENT,
            "mark": (ACCENT, ACCENT_LIGHT, INK, ACCENT),
        },
        {
            "label": "API + BD",
            "tech": "FastAPI · SQLite · JWT",
            "fill": ACCENT_SOFT, "line": None, "fg": INK, "label_color": ACCENT,
            "mark": (ACCENT, INK, ACCENT_LIGHT, ACCENT_DARK),
        },
        {
            "label": "MOTOR DE PATRONES",
            "tech": "NumPy · OpenCV\nscikit-learn · ReportLab",
            "fill": INK, "line": None, "fg": WHITE, "label_color": ACCENT_LIGHT,
            "mark": (ACCENT, ACCENT_LIGHT, ACCENT_LIGHT, ACCENT),
        },
    ]
    bw, bh = 3.35, 2.4
    xs = [MARGIN, 4.99, 9.13]
    for x, box in zip(xs, boxes):
        add_rect(s, x, 2.5, bw, bh, fill=box["fill"], line=box["line"],
                 radius=0.05, shadow=box["fill"] is WHITE)
        coords = [(0, 0), (1, 0), (0, 1), (1, 1)]
        for (cx, cy), color in zip(coords, box["mark"]):
            add_rect(s, x + 0.3 + cx * 0.15, 2.75 + cy * 0.15, 0.115, 0.115,
                     fill=color)
        add_text(s, box["label"], x + 0.3, 3.15, bw - 0.6, 0.28, size=9.5,
                 color=box["label_color"], bold=True, tracking=1.5)
        add_text(s, box["tech"], x + 0.3, 3.52, bw - 0.6, 1.0,
                 size=17 if box["label"] != "MOTOR DE PATRONES" else 16,
                 color=box["fg"], bold=True, line_spacing=1.12)

    add_arrow(s, 4.31, 3.55, 0.57, 0.3, fill=RGBColor.from_string("E9BEC8"))
    add_arrow(s, 8.45, 3.55, 0.57, 0.3, fill=RGBColor.from_string("E9BEC8"))

    add_text(s, "Un contrato de datos común: PatternData", MARGIN, 5.4,
             CONTENT_W, 0.35, size=13, color=MUTED, align=PP_ALIGN.CENTER)

    add_footer(s, 8)
    add_notes(
        s,
        "≈ 25 s — Técnicamente no hay magia: React y TypeScript en el frontend, "
        "FastAPI y SQLite detrás, y un motor en Python con NumPy, OpenCV y "
        "scikit-learn. El PDF se genera con ReportLab. Todo habla por una API REST "
        "con el mismo contrato de datos, PatternData. Eso nos permitió construir en "
        "paralelo sin pisarnos.",
    )


def slide_team(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "Cómo trabajamos")
    add_title(s, "Cinco personas, un solo contrato")
    add_rule(s)

    centers = [1.55, 4.11, 6.66, 9.22, 11.78]
    add_rect(s, centers[0], 2.94, centers[-1] - centers[0], 0.012, fill=BORDER)
    labels = ["Core", "Motor", "Social", "Editor", "Social UI"]
    for i, (cx, label) in enumerate(zip(centers, labels)):
        add_rect(s, cx - 0.36, 2.58, 0.72, 0.72, fill=ACCENT, radius=0.5)
        add_text(s, str(i + 1), cx - 0.36, 2.58, 0.72, 0.72, size=16,
                 color=WHITE, bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, label, cx - 0.85, 3.45, 1.7, 0.3, size=11.5, color=INK,
                 bold=True, align=PP_ALIGN.CENTER)

    principles = [
        "Contrato de API desde el minuto 0",
        "Un dueño por archivo",
        "Agentes de IA (opencode) en paralelo",
    ]
    cw, gap = 3.72, 0.24
    for i, text in enumerate(principles):
        x = MARGIN + i * (cw + gap)
        add_rect(s, x, 4.45, cw, 1.6, fill=WHITE, line=BORDER, radius=0.04,
                 shadow=True)
        add_pixel_mark(s, x + 0.28, 4.74, cell=0.1, gap=0.03)
        add_text(s, text, x + 0.28, 5.17, cw - 0.56, 0.7, size=16,
                 color=INK, line_spacing=1.12)

    add_footer(s, 9)
    add_notes(
        s,
        "≈ 25 s — Y aquí está la parte de la que estamos más orgullosos: cinco "
        "personas, cinco ramas y un contrato de API cerrado desde el minuto cero. "
        "Cada persona era dueña de sus archivos, así que los merges salieron "
        "limpios. Trabajamos con agentes de IA de opencode: nos repartimos el "
        "producto, no los conflictos de Git.",
    )


def slide_next(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "Próximos pasos")
    add_title(s, "Lo que viene")
    add_rule(s)

    cards = [
        ("01", "Más marcas", "Cuentas e hilos de más fabricantes"),
        ("02", "App móvil", "Cámara integrada"),
        ("03", "Patrones 3D", "HAMA en relieve"),
        ("04", "Tienda de kits", "Materiales exactos"),
    ]
    cw, gap = 2.7, 0.28
    for i, (num, title, body) in enumerate(cards):
        x = MARGIN + i * (cw + gap)
        add_rect(s, x, 2.5, cw, 3.55, fill=WHITE, line=BORDER, radius=0.04,
                 shadow=True)
        add_text(s, num, x + 0.3, 2.78, 0.6, 0.3, size=10.5, color=ACCENT,
                 bold=True, tracking=1.2)
        add_text(s, title, x + 0.3, 3.3, cw - 0.6, 1.0, size=18.5, color=INK,
                 bold=True, line_spacing=1.02)
        add_rect(s, x + 0.3, 4.62, cw - 0.6, 0.009, fill=BORDER)
        add_text(s, body, x + 0.3, 4.86, cw - 0.6, 0.65, size=12, color=MUTED,
                 line_spacing=1.12)

    add_footer(s, 10)
    add_notes(
        s,
        "≈ 20 s — ¿Y después? Más marcas de cuentas e hilos, una app móvil con "
        "cámara, patrones 3D para HAMA y, a más largo plazo, una tienda de kits con "
        "los materiales exactos de cada patrón. Cerrar el círculo entre la idea y "
        "la manualidad.",
    )


def slide_closing(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s)
    add_eyebrow(s, "El próximo patrón es tuyo")

    add_text(s, "Pixel Pattern", MARGIN, 1.45, 7.6, 0.8, size=48, color=INK,
             bold=True)
    add_text(s, "¡Gracias!", MARGIN, 2.18, 7.6, 1.1, size=60, color=ACCENT,
             bold=True)
    add_rect(s, MARGIN + 0.02, 3.68, 1.15, 0.07, fill=ACCENT)
    add_text(s, "Demo y código en el QR.", MARGIN + 0.02, 4.02, 7.0, 0.4,
             size=18, color=INK)
    add_text(s, "[enlace a la demo]", MARGIN + 0.02, 4.5, 7.0, 0.35,
             size=13, color=MUTED)

    add_text(s, "[Nombre del equipo] · [Nombre del hackathon]", MARGIN, 6.35,
             8.0, 0.35, size=11.5, color=MUTED)

    qr_x, qr_y, qr_size = 8.95, 1.62, 3.35
    add_rect(s, qr_x, qr_y, qr_size, qr_size, fill=WHITE, line=ACCENT,
             line_w=1.5, radius=0.05, dash=MSO_LINE_DASH_STYLE.DASH)
    add_text(s, "[QR / enlace\na la demo]", qr_x, qr_y, qr_size, qr_size,
             size=13, color=MUTED, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)
    add_heart(s, 9.72, 5.42, cell=0.17, gap=0.025)

    add_text(s, f"{TOTAL_SLIDES:02d} / {TOTAL_SLIDES}", 10.5, 7.02, 1.98,
             0.25, size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    add_notes(
        s,
        "≈ 15 s — Esto es Pixel Pattern: una forma más rápida de pasar de una foto "
        "a algo que puedes hacer con las manos, y una comunidad para compartirlo. "
        "El código y la demo están en el QR. Gracias, y ahora... ¡a hacer patrones!",
    )


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide_cover(prs)
    slide_problem(prs)
    slide_solution(prs)
    slide_demo_convert(prs)
    slide_engine(prs)
    slide_demo_make(prs)
    slide_community(prs)
    slide_tech(prs)
    slide_team(prs)
    slide_next(prs)
    slide_closing(prs)

    prs.core_properties.title = "Pixel Pattern — Pitch"
    prs.core_properties.subject = "Crea, comparte y haz patrones de HAMA y punto de cruz"
    prs.core_properties.author = "Pixel Pattern"
    prs.core_properties.keywords = "HAMA, punto de cruz, patrones, hackathon"

    prs.save(OUT_PATH)
    return OUT_PATH


def main() -> None:
    path = build()
    size_mb = path.stat().st_size / (1024 * 1024)

    print(f"OK  {path}")
    print(f"    {TOTAL_SLIDES} diapositivas · {size_mb:.2f} MB")
    print()
    print("Capturas:")
    for name in ("feed.png", "crear.png", "foto-a-patron.png", "hacer.png",
                 "buscador.png"):
        state = "OK" if (IMG_DIR / name).exists() else "FALTA (placeholder)"
        print(f"  - {name:<20} {state}")
    print()
    print("Sustituye los textos entre corchetes de la portada y el cierre.")


if __name__ == "__main__":
    main()
